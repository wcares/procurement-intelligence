import duckdb
import pandas as pd

from procintel.core import paths
from procintel.intake import classify, extract, link, process, store
from procintel.jobs import store as jobs_store
from procintel.pipeline import store as pipeline_store

ICR = ("The World Bank\nIMPLEMENTATION COMPLETION AND RESULTS REPORT\nRural Roads Project (P123456)\nRepublic of Kenya\n"
       + "The project rehabilitated feeder roads; procurement of civil works was delayed by re-bidding. " * 20)


def _setup(tmp_path, monkeypatch):
    monkeypatch.setattr(paths, "DB_PATH", tmp_path / "main.duckdb")
    monkeypatch.setattr(paths, "PIPELINE_DB", tmp_path / "pipeline.duckdb")
    monkeypatch.setattr(jobs_store, "DB_PATH", tmp_path / "jobs.duckdb")
    monkeypatch.setattr(paths, "DATA_DIR", tmp_path)                 # uploads go to <data>/uploads (local blob store)
    monkeypatch.delenv("PROCINTEL_DECIDE_API_BASE", raising=False)
    con = duckdb.connect(str(tmp_path / "main.duckdb"))
    con.execute("CREATE TABLE projects AS SELECT * FROM (VALUES ('P123456', 'Rural Roads', 'Republic of Kenya', 'KE'), "
                "('P000002', 'Global', 'World', '1W')) t(project_id, name, country, country_code)")
    docs = pd.DataFrame([{"doc_id": "1", "project_id": "P1", "doc_type": "x", "doc_date": pd.Timestamp("2020-01-01", tz="UTC"),
                          "title": "t", "pdf_url": "u", "txt_url": "u", "text": "old"}])
    con.execute("CREATE TABLE documents AS SELECT * FROM docs")
    con.close()


def test_a_completion_report_is_extracted_classified_linked_added_and_queued(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    up = store.save("report.txt", ICR.encode(), "txt", "u1")
    assert up["status"] == "received"
    out = process.run(up["upload_id"])
    assert [s["stage"] for s in out["stages"]] == ["received", "extract", "classify", "link", "add"]
    assert out["status"] == "added" and out["doc_type"] == "completion_report" and out["doc_id"] == f"upload-{up['upload_id']}"
    assert out["links"]["projects"] == [{"id": "P123456", "name": "Rural Roads", "country": "Republic of Kenya", "known": True}]
    assert out["links"]["countries"] == ["Republic of Kenya"]
    con = duckdb.connect(str(tmp_path / "main.duckdb"))
    assert con.execute("SELECT project_id, txt_url FROM documents WHERE doc_id = ?", [out["doc_id"]]).fetchone() == ("P123456", None)
    con.close()
    job = jobs_store.next_pending()
    assert job["kind"] == "pipeline_batch" and job["priority"] == 0         # high priority
    assert {f["rule"] for f in pipeline_store.overview()["firings"]} == {"uploads-add", "watch-alerts"}
    again = store.save("copy.txt", ICR.encode(), "txt", "u1")               # same content: a duplicate, not reprocessed
    assert again["status"] == "duplicate" and up["upload_id"] in again["stages"][0]["detail"]
    assert process.run(again["upload_id"])["status"] == "duplicate"


def test_other_kinds_are_kept_and_unreadable_or_unlinked_ones_go_to_review(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    notice = "INVITATION FOR BIDS\nRoads in Kenya, project P999999\n" + "Bidders shall submit sealed bids. " * 30
    out = process.run(store.save("ifb.txt", notice.encode(), "txt", None)["upload_id"])
    assert out["status"] == "kept" and out["doc_type"] == "procurement_notice"
    assert out["links"]["projects"][0] == {"id": "P999999", "name": None, "country": None, "known": False}
    assert out["links"]["countries"] == ["Republic of Kenya"]          # found as "Kenya"
    assert process.run(store.save("scan.txt", b"x" * 100, "txt", None)["upload_id"])["status"] == "needs_review"
    isr = ICR.replace("IMPLEMENTATION COMPLETION AND RESULTS REPORT", "Implementation Status & Results Report")
    out = process.run(store.save("isr.txt", isr.encode(), "txt", None)["upload_id"])
    assert out["status"] == "kept" and out["doc_type"] == "supervision_report"
    assert out["stages"][2]["detail"] == "supervision report (from its title)"
    assert out["links"]["countries"] == ["Republic of Kenya"]                # "World" (a region code) is not a country
    unlinked = ICR.replace("P123456", "P000001")
    assert process.run(store.save("icr2.txt", unlinked.encode(), "txt", None)["upload_id"])["status"] == "needs_review"
    assert jobs_store.next_pending() is None                                # nothing was added to the corpus


def test_formats_and_keyword_classes():
    assert extract.format_of("a.PDF") == "pdf" and extract.format_of("a.exe") is None and extract.format_of("noext") is None
    assert classify.by_keywords("Implementation Status and Results Report, Sequence 12") == "supervision_report"
    assert classify.by_keywords("AMENDMENT No. 2 to the contract") == "amendment"
    assert classify.by_keywords("minutes of a meeting") == "other"
    # the earliest title phrase wins: an ISR names the completion report it leads to
    assert classify.by_keywords("Implementation Status & Results Report ... the Implementation Completion and Results "
                                "Report will follow") == "supervision_report"
    assert link.PROJECT_ID.findall("P123456, P1234567 and xP123456") == ["P123456"]


def test_docx_and_html_text(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Contract agreement for works")
    d.add_table(rows=1, cols=2).rows[0].cells[0].text = "Lot 1"
    d.save(tmp_path / "c.docx")
    text = extract.pages(tmp_path / "c.docx", "docx")[0]
    assert "Contract agreement" in text and "Lot 1" in text
    (tmp_path / "p.html").write_text("<html><script>x()</script><p>Request for bids</p></html>")
    assert extract.pages(tmp_path / "p.html", "html")[0].strip() == "Request for bids"


def test_upload_endpoint_checks_type_and_size_and_processes_in_the_background(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from procintel.api.app import app
    _setup(tmp_path, monkeypatch)
    monkeypatch.setattr(extract, "MAX_BYTES", 10_000)
    c = TestClient(app)
    assert c.post("/intake/uploads", files={"file": ("x.exe", b"MZ")}).status_code == 400
    assert c.post("/intake/uploads", files={"file": ("big.txt", b"a" * 10_001)}).status_code == 413
    r = c.post("/intake/uploads", files={"file": ("icr.txt", ICR.encode(), "text/plain")})
    assert r.status_code == 200 and r.json()["status"] == "received"
    got = c.get(f"/intake/uploads/{r.json()['upload_id']}").json()   # the background task ran after the response
    assert got["status"] == "added" and c.get("/intake/uploads").json()["uploads"][0]["upload_id"] == got["upload_id"]
    assert c.get("/intake/uploads/nope").status_code == 404


def test_a_refresh_keeping_existing_reports_keeps_uploaded_text(tmp_path, monkeypatch):
    from procintel.ingestion import ingest
    db_file = tmp_path / "main.duckdb"
    monkeypatch.setattr(ingest, "DB_PATH", db_file)
    docs = pd.DataFrame([
        {"doc_id": "1", "project_id": "P1", "doc_type": "x", "doc_date": pd.Timestamp("2020-01-01", tz="UTC"),
         "title": "t", "pdf_url": "u", "txt_url": "http://wb/1.txt", "text": "stale"},
        {"doc_id": "upload-abc", "project_id": "P1", "doc_type": "x", "doc_date": pd.Timestamp("2026-10-25", tz="UTC"),
         "title": "mine", "pdf_url": None, "txt_url": None, "text": "uploaded text"}])
    con = duckdb.connect(str(db_file))
    con.execute("CREATE TABLE documents AS SELECT * FROM docs")
    con.close()

    class Client:
        def __init__(self, **kw): pass
        def iter_documents(self, n, **kw): return iter([])
        def projects(self, rows, id): return 1, [{"id": id, "project_name": "p"}]
        def iter_notices(self, n, project_id): return iter([{}])
        def document_text(self, url):
            assert url, "an uploaded report must not be fetched"
            return "fresh"
    monkeypatch.setattr(ingest, "WBClient", Client)
    monkeypatch.setattr(ingest, "notice_row", lambda n: {"notice_id": "n1"})  # an empty table cannot be written
    ingest.run(5, keep_existing=True)
    con = duckdb.connect(str(db_file))
    assert dict(con.execute("SELECT doc_id, text FROM documents").fetchall()) == {"1": "fresh", "upload-abc": "uploaded text"}
    con.close()


def test_the_title_beats_the_model_and_a_model_only_completion_report_waits_for_review(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    said = {"completion_report": 0.84, "supervision_report": 0.1, "procurement_notice": 0.02, "contract": 0.02,
            "amendment": 0.01, "other": 0.01}
    monkeypatch.setattr(classify, "_model", lambda text: said)       # what the model said about a real ISR
    isr = ICR.replace("IMPLEMENTATION COMPLETION AND RESULTS REPORT", "Implementation Status & Results Report")
    k = classify.classify(isr)
    assert k["type"] == "supervision_report" and k["titled"] and "decision model said completion report (84%)" in k["method"]
    untitled = ICR.replace("IMPLEMENTATION COMPLETION AND RESULTS REPORT", "Project evaluation")
    out = process.run(store.save("eval.txt", untitled.encode(), "txt", None)["upload_id"])
    assert out["doc_type"] == "completion_report" and out["status"] == "needs_review"
    assert jobs_store.next_pending() is None                                 # not added to the evidence


def test_removing_an_upload_deletes_its_file_and_takes_its_report_out_of_the_evidence(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    up = process.run(store.save("report.txt", ICR.encode(), "txt", "u1")["upload_id"])
    stored = tmp_path / up["path"]
    assert stored.exists() and up["status"] == "added"
    while jobs_store.next_pending():                                  # the add's job; this test looks at the removal's
        pass
    out = process.remove(up["upload_id"])
    assert out["status"] == "removed" and out["path"] is None and not stored.exists()
    assert out["stages"][-1]["stage"] == "removed" and "taken out of the evidence" in out["stages"][-1]["detail"]
    con = duckdb.connect(str(tmp_path / "main.duckdb"))
    assert con.execute("SELECT count(*) FROM documents WHERE doc_id = ?", [up["doc_id"]]).fetchone()[0] == 0
    assert con.execute("SELECT count(*) FROM documents").fetchone()[0] == 1       # the other report stays
    con.close()
    pcon = pipeline_store.connect()
    assert pcon.execute("SELECT type, key FROM event_log WHERE type = 'Upload.removed'").fetchall() == [
        ("Upload.removed", up["doc_id"])]
    pcon.close()
    assert jobs_store.next_pending() is not None                                  # re-cut and re-index queued
    assert process.remove(up["upload_id"])["status"] == "removed"                 # a second removal changes nothing
    assert process.remove("nope") is None
