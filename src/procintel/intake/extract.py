"""Text by format: one string per page (PDF) or one for the whole file. Adding a format = one reader in READERS.

PDF pages without a text layer (scans) are read by OCR (procintel.intake.ocr) when an engine is available;
pages_ocr() also says how many pages that was. Without OCR such a file is reported for review rather than guessed at."""
from __future__ import annotations

from pathlib import Path
from typing import Callable

MAX_BYTES = 50 * 1024 * 1024
SCANNED_CHARS = 40    # a PDF page with less text than this is treated as a scan


def _pdf(path: Path) -> list[str]:
    return _pdf_ocr(path)[0]


def _pdf_ocr(path: Path) -> tuple[list[str], int]:
    from pypdf import PdfReader
    from procintel.intake import ocr
    pages = [p.extract_text() or "" for p in PdfReader(str(path)).pages]
    scans = [i for i, t in enumerate(pages) if len(t.strip()) < SCANNED_CHARS]
    read = ocr.read_pdf_pages(path, scans)
    for i, t in read.items():
        pages[i] = t
    return pages, len(read)


def _docx(path: Path) -> list[str]:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
    return ["\n".join(parts)]


def _text(path: Path) -> list[str]:
    return [path.read_text(errors="replace")]


def _html(path: Path) -> list[str]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_text(errors="replace"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return [soup.get_text("\n")]


READERS: dict[str, Callable[[Path], list[str]]] = {
    "pdf": _pdf, "docx": _docx, "txt": _text, "md": _text, "html": _html, "htm": _html}


def format_of(filename: str) -> str | None:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext if ext in READERS else None


def pages(path: Path, fmt: str) -> list[str]:
    return READERS[fmt](path)


def pages_ocr(path: Path, fmt: str) -> tuple[list[str], int]:
    """(pages, how many of them were read by OCR)."""
    return _pdf_ocr(path) if fmt == "pdf" else (pages(path, fmt), 0)
